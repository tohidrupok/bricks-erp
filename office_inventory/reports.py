"""
Reporting - Excel (.xlsx) and PDF export for every major section
(Inventory, Requisitions, Purchases, Assignments, Damage records).

Two generic helpers do all the heavy lifting:
    build_excel(title, headers, rows)  -> HttpResponse (.xlsx)
    build_pdf(title, headers, rows)    -> HttpResponse (.pdf)

Each `<module>_report_rows()` function below just turns a queryset into a
flat list of (headers, rows) ready to hand to either helper, so views.py
stays thin: fetch filtered queryset -> build rows -> build_excel/build_pdf.
"""
from datetime import datetime

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


BRAND_HEX = "0D3B66"


# ---------------------------------------------------------------------------
# Generic export helpers
# ---------------------------------------------------------------------------

def build_excel(title, headers, rows, filename=None):
    """headers: list[str]; rows: list[list] (same order as headers)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    ws.cell(row=1, column=1).font = Font(size=14, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor=BRAND_HEX)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.append([f"Generated: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B5FA8")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append(list(row))

    # Auto width
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in rows:
            val = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    filename = filename or f"{title.lower().replace(' ', '_')}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def build_pdf(title, headers, rows, filename=None, landscape_mode=True, subtitle=None):
    filename = filename or f"{title.lower().replace(' ', '_')}.pdf"
    response = HttpResponse(content_type="application/pdf")
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    pagesize = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(
        response, pagesize=pagesize,
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    elements = []

    title_style = styles['Title']
    title_style.textColor = colors.HexColor(f"#{BRAND_HEX}")
    elements.append(Paragraph(title, title_style))
    meta = subtitle or f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    elements.append(Paragraph(meta, styles['Normal']))
    elements.append(Spacer(1, 10))

    data = [headers] + [[str(c) if c is not None else "-" for c in row] for row in rows]
    if len(rows) == 0:
        data.append(["No data found for the selected filters." if headers else ""] +
                     [""] * (len(headers) - 1))

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_HEX}")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    doc.build(elements)
    return response


# ---------------------------------------------------------------------------
# Per-module row builders - each returns (headers, rows)
# ---------------------------------------------------------------------------

def items_report_rows(items):
    headers = ["Code", "Name", "Category", "Location", "Unit", "Min Stock", "Current Stock", "Status", "Is Asset"]
    rows = [[
        i.item_code, i.name,
        i.category.name if i.category_id else "-",
        i.location.name if i.location_id else "-",
        i.get_unit_display(), i.min_stock, i.current_stock, i.status,
        "Yes" if i.is_asset else "No",
    ] for i in items]
    return headers, rows


def requisitions_report_rows(requisitions):
    headers = ["Req No", "Date", "Department", "Employee", "Item", "Requested",
               "Approved", "Issued", "Status"]
    rows = [[
        r.req_no, r.date, r.department.name, str(r.employee), str(r.item),
        r.requested_qty, r.approved_qty or "-", r.issued_qty or "-",
        r.get_status_display(),
    ] for r in requisitions]
    return headers, rows


def purchases_report_rows(purchases):
    headers = ["Purchase No", "Date", "Item", "Supplier", "Qty", "Unit Price",
               "Total Amount", "Status", "Received By"]
    rows = [[
        p.purchase_no, p.date, str(p.item), p.supplier or "-", p.qty,
        p.unit_price, p.total_amount, p.get_status_display(),
        str(p.received_by) if p.received_by_id else "-",
    ] for p in purchases]
    return headers, rows


def assignments_report_rows(assignments):
    headers = ["Assign Type", "Assigned To", "Location", "Item", "Qty",
               "Assigned Date", "Return Date", "Status", "Remarks"]
    rows = [[
        a.assignment_type.name if a.assignment_type_id else "-",
        a.assignee_label,
        a.location.name if a.location_id else "-",
        str(a.item), a.qty, a.assigned_date, a.return_date or "-",
        a.get_status_display(), a.remarks or "-",
    ] for a in assignments]
    return headers, rows


def damage_report_rows(records):
    headers = ["Damage No", "Date", "Item", "Qty", "Responsible Type",
               "Responsible", "Reason", "Severity", "Status", "Est. Cost",
               "Stock Deducted"]
    rows = [[
        d.damage_no, d.date_reported, str(d.item), d.qty,
        d.responsible_type.name if d.responsible_type_id else "-",
        d.responsible_label, d.reason.name if d.reason_id else "-",
        d.get_severity_display(), d.get_status_display(), d.estimated_cost,
        "Yes" if d.deduct_from_stock else "No",
    ] for d in records]
    return headers, rows
